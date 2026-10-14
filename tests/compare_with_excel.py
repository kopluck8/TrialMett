#!/usr/bin/env python3
"""
Uji kesetaraan: hasil mesin rumus web vs hasil perhitungan spreadsheet (LibreOffice)
atas FILE EXCEL ASLI dengan isian yang sama.

    node tests/make_scenarios.js tests/out
    python3 tests/compare_with_excel.py "METT 4 _Indonesia_KSDAE_VER4.4.xlsm" tests/out

Langkah per skenario:
  1. salin file .xlsm asli, isi sel-sel isian (alamat sel sama persis) dengan openpyxl
  2. hitung ulang semua rumus dengan LibreOffice (headless)
  3. bandingkan SETIAP sel rumus dengan nilai dari mesin web (tests/out/expected.json)

Butuh: openpyxl, LibreOffice (soffice).
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import warnings

import openpyxl

warnings.filterwarnings("ignore")

# LibreOffice di server memakai lokal Inggris untuk TEXT(..,"dddd"/"mmmm");
# aplikasi web memakai nama Indonesia (seperti Excel berlokal Indonesia).
EN_ID = {"Sunday": "Minggu", "Monday": "Senin", "Tuesday": "Selasa", "Wednesday": "Rabu", "Thursday": "Kamis",
         "Friday": "Jumat", "Saturday": "Sabtu", "January": "Januari", "February": "Februari", "March": "Maret",
         "April": "April", "May": "Mei", "June": "Juni", "July": "Juli", "August": "Agustus",
         "September": "September", "October": "Oktober", "November": "November", "December": "Desember"}
EXCEL_EPOCH = None


def norm(v):
    import datetime
    if isinstance(v, datetime.datetime):
        d = v - datetime.datetime(1899, 12, 30)
        return d.days + d.seconds / 86400
    if isinstance(v, datetime.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, str):
        for en, idn in EN_ID.items():
            v = v.replace(en, idn)
        return v
    return v


# Perbedaan yang berasal dari LibreOffice, bukan dari aplikasi web. Keduanya sudah dicek
# terhadap nilai cache Excel asli (tests/check_excel_cache.js): web = Excel.
KNOWN_LIBREOFFICE = {
    ("PA attributes data sheet", "CU2"): "Rumus ='Atribut Kawasan'!#REF!: Excel -> #REF!, LibreOffice -> #VALUE!",
    ("Print BA", "B14"): "Tanggal kosong (serial 0): Excel -> '00 Januari 1900', LibreOffice -> '30 Desember 1899'",
}


def same(exp, got):
    if isinstance(exp, dict) and "error" in exp:
        return isinstance(got, str) and got.strip() == exp["error"]
    if got is None:
        return exp in (0, "", False, None)
    if isinstance(exp, bool) or isinstance(got, bool):
        return bool(exp) == bool(got) and type(exp) == type(got)
    if isinstance(exp, (int, float)) and isinstance(got, (int, float)):
        return abs(exp - got) < 1e-9 * max(1, abs(exp))
    return str(exp) == str(got)


def main(src, outdir):
    scenarios = json.load(open(os.path.join(outdir, "scenarios.json"), encoding="utf-8"))
    expected = json.load(open(os.path.join(outdir, "expected.json"), encoding="utf-8"))
    tmp = tempfile.mkdtemp(prefix="mett_cmp_")
    total_cells = total_diff = 0
    report = []
    known_hits = set()
    for sc in scenarios:
        name = sc["name"]
        wb = openpyxl.load_workbook(src, keep_vba=True)
        for key, val in sc["inputs"].items():
            sheet, addr = key.rsplit("!", 1)
            if sheet not in wb.sheetnames:
                continue
            wb[sheet][addr].value = val
        f_in = os.path.join(tmp, f"{name}.xlsm")
        wb.save(f_in)
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, f_in],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
        calc = openpyxl.load_workbook(os.path.join(tmp, f"{name}.xlsx"), data_only=True)
        diffs = []
        n = 0
        for sheet, cells in expected[name].items():
            ws = calc[sheet]
            for addr, exp in cells.items():
                n += 1
                got = norm(ws[addr].value)
                if not same(exp, got):
                    known = KNOWN_LIBREOFFICE.get((sheet, addr))
                    if known and (addr != "B14" or "1899" in str(got)):
                        known_hits.add(known)
                        continue
                    diffs.append((sheet, addr, exp, got))
        total_cells += n
        total_diff += len(diffs)
        report.append(f"{name:10s}: {n} sel rumus dibandingkan, {len(diffs)} berbeda")
        for d in diffs[:15]:
            report.append(f"    BEDA {d[0]}!{d[1]}  web={d[2]!r}  libreoffice={d[3]!r}")
    shutil.rmtree(tmp, ignore_errors=True)
    print("\n".join(report))
    for k in sorted(known_hits):
        print("  (diabaikan, perbedaan LibreOffice vs Excel) " + k)
    print(f"TOTAL: {total_cells} perbandingan, {total_diff} berbeda")
    return 0 if total_diff == 0 else 1


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2]))
